"""
文件解析：把上传的二进制文件变成纯文本。

支持 pdf / docx / md / txt 四类。解析器依赖（pypdf、python-docx）
全部延迟导入：没装依赖时本模块仍可导入，只有真正解析对应格式才报错，
这样"知识库功能关掉时不影响整个应用启动"。
"""

import io
from pathlib import Path

# 白名单之外的扩展名一律拒绝，不做"猜格式"的兜底
SUPPORTED_EXTENSIONS: dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".md": "md",
    ".markdown": "md",
    ".txt": "txt",
}

# 前端提示用
SUPPORTED_LABEL = "PDF、Word(.docx)、Markdown、纯文本(.txt)"

# 按顺序尝试的文本编码。gb18030 覆盖 GBK/GB2312，
# 中文 Windows 上导出的 txt 大多是这个编码
_TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "utf-16", "big5")


class KbParseError(Exception):
    """解析失败。错误信息会直接展示给用户，必须可读。"""


def detect_ext(filename: str) -> str:
    """从文件名取扩展名（可能为空串，由调用方负责拒绝）。"""
    return Path(filename or "").suffix.lower()


def is_supported(filename: str) -> bool:
    return detect_ext(filename) in SUPPORTED_EXTENSIONS


def parse(filename: str, data: bytes) -> tuple[str, str]:
    """
    解析文件，返回 (扩展名, 纯文本)。

    扩展名不含点，例如 "pdf"；前端按它给文档分组。
    """
    if not data:
        raise KbParseError("文件内容为空")

    suffix = detect_ext(filename)
    kind = SUPPORTED_EXTENSIONS.get(suffix)
    if kind is None:
        raise KbParseError(f"不支持的文件类型：{suffix or '未知'}。目前支持 {SUPPORTED_LABEL}")

    if kind in ("md", "txt"):
        text = _decode_text(data)
    elif kind == "pdf":
        text = _parse_pdf(data)
    else:
        text = _parse_docx(data)

    text = text.strip()
    if not text:
        raise KbParseError("没有解析出任何文本内容")
    return kind, text


def _decode_text(data: bytes) -> str:
    """按常见编码逐个尝试解码纯文本。"""
    for encoding in _TEXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    # 全部失败时用替换字符兜底，宁可有几个乱码也不要整体失败
    return data.decode("utf-8", errors="replace")


def _parse_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover - 依赖缺失时的兜底
        raise KbParseError("缺少 pypdf 依赖，无法解析 PDF") from e

    try:
        reader = PdfReader(io.BytesIO(data))
        if getattr(reader, "is_encrypted", False):
            # 空密码能解开的加密 PDF 很常见，先试一下
            try:
                reader.decrypt("")
            except Exception as e:  # noqa: BLE001
                raise KbParseError("该 PDF 已加密，无法解析") from e

        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except KbParseError:
        raise
    except Exception as e:  # noqa: BLE001
        raise KbParseError(f"PDF 解析失败：{e}") from e

    text = "\n\n".join(p for p in pages if p)
    if not text.strip():
        # 扫描件没有文本层，属于常见情况，给出可操作的提示
        raise KbParseError("这个 PDF 里没有可提取的文字（可能是扫描件或纯图片）")
    return text


def _parse_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as e:  # pragma: no cover - 依赖缺失时的兜底
        raise KbParseError("缺少 python-docx 依赖，无法解析 Word 文档") from e

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise KbParseError(f"Word 文档解析失败：{e}") from e

    lines = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]

    # 表格内容也要收进来，否则合同、参数表这类文档会丢一大块
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = "\t".join(c for c in cells if c)
            if line:
                lines.append(line)

    text = "\n".join(lines)
    if not text.strip():
        raise KbParseError("Word 文档里没有正文内容")
    return text
